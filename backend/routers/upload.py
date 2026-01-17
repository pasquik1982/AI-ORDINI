import logging
import io
import base64
from datetime import datetime
from typing import List, Dict, Any, Optional
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy.ext.asyncio import AsyncSession
from pydantic import BaseModel

from core.database import get_db
from dependencies.auth import get_current_user
from schemas.auth import UserResponse
from services.acquisti import AcquistiService
from services.vendite import VenditeService
from services.fornitori import FornitoriService

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/api/v1/upload", tags=["upload"])


class ParsedSupplierRow(BaseModel):
    fornitore: str
    data_ordine: str
    data_consegna: str
    categoria_prodotto: str
    quantita: float
    prezzo_totale: float


class ParsedSalesRow(BaseModel):
    data: str
    canale_vendita: str
    categoria_prodotto: str
    quantita: float
    prezzo_totale: Optional[float] = None


class FileParseRequest(BaseModel):
    file_content: str  # Base64 encoded file content
    file_name: str


class ParseSupplierResponse(BaseModel):
    rows: List[ParsedSupplierRow]
    total_rows: int


class ParseSalesResponse(BaseModel):
    rows: List[ParsedSalesRow]
    total_rows: int


class ImportSupplierRequest(BaseModel):
    rows: List[ParsedSupplierRow]


class ImportSalesRequest(BaseModel):
    rows: List[ParsedSalesRow]


class ImportResponse(BaseModel):
    success: bool
    imported_count: int
    message: str


def parse_excel_date(cell_value) -> Optional[datetime]:
    """Parse Excel date from cell value (handles both datetime and serial number)"""
    if cell_value is None:
        return None
    
    # If it's already a datetime object
    if isinstance(cell_value, datetime):
        return cell_value
    
    # If it's a date object
    from datetime import date
    if isinstance(cell_value, date):
        return datetime.combine(cell_value, datetime.min.time())
    
    # If it's a number (Excel serial date)
    if isinstance(cell_value, (int, float)):
        try:
            # Excel dates start from 1900-01-01 (serial 1)
            from datetime import timedelta
            base_date = datetime(1899, 12, 30)  # Excel's epoch
            return base_date + timedelta(days=cell_value)
        except:
            pass
    
    # If it's a string, try to parse it
    if isinstance(cell_value, str):
        return parse_date_string(cell_value)
    
    return None


def parse_date_string(date_str: str) -> Optional[datetime]:
    """Parse various date string formats and return datetime object"""
    if not date_str or not isinstance(date_str, str) or date_str.strip() == "":
        return None
    
    date_str = date_str.strip()
    
    # Try common date formats
    formats = [
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%Y-%m-%d",
        "%d/%m/%y",
        "%d.%m.%Y",
        "%Y/%m/%d",
        "%d/%m/%Y %H:%M:%S",
        "%Y-%m-%d %H:%M:%S"
    ]
    
    for fmt in formats:
        try:
            return datetime.strptime(date_str, fmt)
        except ValueError:
            continue
    
    logger.warning(f"Could not parse date string: {date_str}")
    return None


def parse_date(date_value) -> str:
    """Parse date from various formats and return ISO format string"""
    # Try Excel date parsing first
    dt = parse_excel_date(date_value)
    
    if dt:
        return dt.strftime("%Y-%m-%d %H:%M:%S")
    
    # If parsing failed, log warning and return empty string
    logger.warning(f"Could not parse date value: {date_value}, type: {type(date_value)}")
    return ""


def parse_date_to_datetime(date_value) -> Optional[datetime]:
    """Parse date from various formats and return datetime object"""
    # If it's already a datetime
    if isinstance(date_value, datetime):
        return date_value
    
    # If it's a string in ISO format
    if isinstance(date_value, str):
        if date_value.strip() == "":
            return None
        # Try to parse ISO format first
        try:
            return datetime.fromisoformat(date_value.replace("Z", "+00:00"))
        except:
            pass
    
    # Try Excel date parsing
    dt = parse_excel_date(date_value)
    
    if dt:
        return dt
    
    logger.warning(f"Could not parse date value: {date_value}, type: {type(date_value)}")
    return None


def parse_number(value) -> float:
    """Parse number from various formats"""
    if value is None or (isinstance(value, str) and value.strip() == ""):
        return 0.0
    
    # If it's already a number
    if isinstance(value, (int, float)):
        return float(value)
    
    # If it's a string, clean and parse
    if isinstance(value, str):
        # Remove common currency symbols and spaces
        cleaned = value.strip().replace("€", "").replace("$", "").replace(",", ".").replace(" ", "")
        try:
            return float(cleaned)
        except ValueError:
            logger.warning(f"Could not parse number: {value}, using 0.0")
            return 0.0
    
    return 0.0


def identify_columns(headers: List[str]) -> Dict[str, int]:
    """Identify column indices based on header names"""
    column_map = {}
    
    # Normalize headers to lowercase for matching
    normalized_headers = [h.lower().strip() if h else "" for h in headers]
    
    # Define possible column name variations
    supplier_names = ["fornitore", "supplier", "vendor"]
    order_date_names = ["data ordine", "data_ordine", "order date", "date"]
    delivery_date_names = ["data consegna", "data_consegna", "delivery date", "consegna"]
    category_names = ["categoria", "category", "prodotto", "product"]
    quantity_names = ["quantita", "quantità", "quantity", "qty", "qta"]
    price_names = ["prezzo", "price", "totale", "total", "importo", "amount"]
    channel_names = ["canale", "channel", "canale vendita"]
    
    for idx, header in enumerate(normalized_headers):
        if any(name in header for name in supplier_names):
            column_map["fornitore"] = idx
        elif any(name in header for name in order_date_names):
            column_map["data_ordine"] = idx
        elif any(name in header for name in delivery_date_names):
            column_map["data_consegna"] = idx
        elif any(name in header for name in category_names):
            column_map["categoria"] = idx
        elif any(name in header for name in quantity_names):
            column_map["quantita"] = idx
        elif any(name in header for name in price_names):
            column_map["prezzo"] = idx
        elif any(name in header for name in channel_names):
            column_map["canale"] = idx
    
    return column_map


def parse_sales_matrix(sheet) -> List[ParsedSalesRow]:
    """Parse sales data from matrix format where:
    - Column A: Sales channel (Provenienza vendite)
    - Columns B-G: Product categories (6 categories)
    - Column H: Date (Mese)
    - Cell values: Quantities for each channel-category combination
    """
    rows = []
    
    # Get headers from first row (categories are in columns B-G)
    first_row = list(sheet[1])
    categories = []
    date_col_idx = 7  # Column H (index 7)
    
    # Identify category columns (B-G, indices 1-6)
    for idx in range(1, 7):
        if idx < len(first_row) and first_row[idx].value:
            categories.append((idx, str(first_row[idx].value).strip()))
    
    # If no categories found in header, use default
    if not categories:
        logger.warning("No category headers found, using default column positions B-G")
        categories = [(i, f"Categoria_{i-1}") for i in range(1, 7)]
    
    logger.info(f"Found categories: {categories}")
    
    # Parse data rows (starting from row 2)
    for row_idx, row_cells in enumerate(sheet.iter_rows(min_row=2), start=2):
        # Get row values
        row = [cell.value for cell in row_cells]
        
        if not any(row):  # Skip empty rows
            continue
        
        # Column A: Sales channel
        canale_vendita = str(row[0]).strip() if row[0] else ""
        if not canale_vendita or canale_vendita.lower() == "none":
            continue
        
        # Column H (index 7): Date
        data_str = ""
        if len(row) > date_col_idx and row[date_col_idx]:
            # Parse the date from column H
            parsed_date = parse_excel_date(row[date_col_idx])
            if parsed_date:
                data_str = parsed_date.strftime("%Y-%m-%d %H:%M:%S")
                logger.info(f"Row {row_idx}: Parsed date from column H: {row[date_col_idx]} -> {data_str}")
            else:
                logger.warning(f"Row {row_idx}: Could not parse date from column H: {row[date_col_idx]}")
        
        if not data_str:
            logger.warning(f"Row {row_idx}: No valid date found in column H, skipping row")
            continue
        
        # Process each category column (B-G)
        for col_idx, categoria_prodotto in categories:
            if len(row) > col_idx and row[col_idx]:
                quantita = parse_number(row[col_idx])
                
                # Only add row if quantity > 0
                if quantita > 0:
                    parsed_row = ParsedSalesRow(
                        data=data_str,
                        canale_vendita=canale_vendita,
                        categoria_prodotto=categoria_prodotto,
                        quantita=quantita,
                        prezzo_totale=None
                    )
                    rows.append(parsed_row)
    
    logger.info(f"Parsed {len(rows)} sales records from matrix")
    return rows


@router.post("/parse-suppliers", response_model=ParseSupplierResponse)
async def parse_supplier_file(
    data: FileParseRequest,
    current_user: UserResponse = Depends(get_current_user),
):
    """Parse supplier file (CSV or Excel) and return preview data"""
    try:
        # Decode base64 file content
        file_content = base64.b64decode(data.file_content)
        
        rows = []
        
        # Determine file type and parse accordingly
        if data.file_name.endswith(('.xlsx', '.xls')):
            # Parse Excel file
            import openpyxl
            workbook = openpyxl.load_workbook(io.BytesIO(file_content))
            sheet = workbook.active
            
            # Get headers from first row
            headers = [cell.value for cell in sheet[1]]
            column_map = identify_columns(headers)
            
            logger.info(f"Column mapping: {column_map}")
            
            # Parse data rows
            for row_idx, row_cells in enumerate(sheet.iter_rows(min_row=2), start=2):
                row = [cell.value for cell in row_cells]
                
                if not any(row):  # Skip empty rows
                    continue
                
                # Get data_consegna from the file
                data_consegna_value = row[column_map.get("data_consegna", 2)] if len(row) > column_map.get("data_consegna", 2) else None
                data_consegna_str = parse_date(data_consegna_value)
                
                if not data_consegna_str:
                    logger.warning(f"Row {row_idx}: Missing or invalid data_consegna, skipping")
                    continue
                
                # Get data_ordine from the file
                data_ordine_value = row[column_map.get("data_ordine", 1)] if len(row) > column_map.get("data_ordine", 1) else None
                data_ordine_str = parse_date(data_ordine_value)
                
                if not data_ordine_str:
                    # If data_ordine is missing, use data_consegna
                    data_ordine_str = data_consegna_str
                
                parsed_row = ParsedSupplierRow(
                    fornitore=str(row[column_map.get("fornitore", 0)] or ""),
                    data_ordine=data_ordine_str,
                    data_consegna=data_consegna_str,
                    categoria_prodotto=str(row[column_map.get("categoria", 3)] or ""),
                    quantita=parse_number(row[column_map.get("quantita", 4)] if len(row) > column_map.get("quantita", 4) else 0),
                    prezzo_totale=parse_number(row[column_map.get("prezzo", 5)] if len(row) > column_map.get("prezzo", 5) else 0)
                )
                rows.append(parsed_row)
        
        else:
            # Parse CSV file
            import csv
            content = file_content.decode('utf-8-sig')
            csv_reader = csv.reader(io.StringIO(content))
            
            # Get headers
            headers = next(csv_reader)
            column_map = identify_columns(headers)
            
            # Parse data rows
            for row in csv_reader:
                if not any(row):  # Skip empty rows
                    continue
                
                data_consegna_value = row[column_map.get("data_consegna", 2)] if len(row) > column_map.get("data_consegna", 2) else ""
                data_consegna_str = parse_date(data_consegna_value)
                
                if not data_consegna_str:
                    logger.warning(f"Missing or invalid data_consegna, skipping row")
                    continue
                
                data_ordine_value = row[column_map.get("data_ordine", 1)] if len(row) > column_map.get("data_ordine", 1) else ""
                data_ordine_str = parse_date(data_ordine_value)
                
                if not data_ordine_str:
                    data_ordine_str = data_consegna_str
                
                parsed_row = ParsedSupplierRow(
                    fornitore=row[column_map.get("fornitore", 0)] if len(row) > column_map.get("fornitore", 0) else "",
                    data_ordine=data_ordine_str,
                    data_consegna=data_consegna_str,
                    categoria_prodotto=row[column_map.get("categoria", 3)] if len(row) > column_map.get("categoria", 3) else "",
                    quantita=parse_number(row[column_map.get("quantita", 4)] if len(row) > column_map.get("quantita", 4) else "0"),
                    prezzo_totale=parse_number(row[column_map.get("prezzo", 5)] if len(row) > column_map.get("prezzo", 5) else "0")
                )
                rows.append(parsed_row)
        
        logger.info(f"Parsed {len(rows)} supplier rows")
        return ParseSupplierResponse(rows=rows, total_rows=len(rows))
    
    except Exception as e:
        logger.error(f"Error parsing supplier file: {e}", exc_info=True)
        raise HTTPException(status_code=400, detail=f"Errore nel parsing del file: {str(e)}")


@router.post("/import-suppliers", response_model=ImportResponse)
async def import_supplier_data(
    data: ImportSupplierRequest,
    current_user: UserResponse = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Import parsed supplier data into database"""
    try:
        acquisti_service = AcquistiService(db)
        fornitori_service = FornitoriService(db)
        
        imported_count = 0
        
        for row in data.rows:
            # Find or create supplier
            supplier = await fornitori_service.get_by_field("nome", row.fornitore)
            if not supplier:
                supplier = await fornitori_service.create({
                    "nome": row.fornitore,
                    "created_at": datetime.now()
                })
            
            # Parse dates
            data_ordine_dt = parse_date_to_datetime(row.data_ordine)
            data_consegna_dt = parse_date_to_datetime(row.data_consegna)
            
            if not data_consegna_dt:
                logger.warning(f"Skipping row with invalid data_consegna: {row.data_consegna}")
                continue
            
            if not data_ordine_dt:
                data_ordine_dt = data_consegna_dt
            
            # Create purchase record
            await acquisti_service.create({
                "fornitore_id": supplier.id,
                "data_ordine": data_ordine_dt,
                "data_consegna": data_consegna_dt,
                "categoria_prodotto": row.categoria_prodotto,
                "quantita": row.quantita,
                "prezzo_totale": row.prezzo_totale,
                "created_at": datetime.now()
            }, user_id=current_user.id)
            
            imported_count += 1
        
        return ImportResponse(
            success=True,
            imported_count=imported_count,
            message=f"Importati con successo {imported_count} ordini"
        )
    
    except Exception as e:
        logger.error(f"Error importing supplier data: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"Errore durante l'importazione: {str(e)}")


@router.post("/parse-sales", response_model=ParseSalesResponse)
async def parse_sales_file(
    data: FileParseRequest,
    current_user: UserResponse = Depends(get_current_user),
):
    """Parse sales file in matrix format where:
    - Column A: Sales channel
    - Columns B-G: Product categories
    - Column H: Date (Mese)
    - Cell values: Quantities
    """
    try:
        # Decode base64 file content
        file_content = base64.b64decode(data.file_content)
        
        rows = []
        
        # Determine file type and parse accordingly
        if data.file_name.endswith(('.xlsx', '.xls')):
            # Parse Excel file with matrix format
            import openpyxl
            workbook = openpyxl.load_workbook(io.BytesIO(file_content))
            
            # Process all sheets
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                logger.info(f"Processing sheet: {sheet_name}")
                sheet_rows = parse_sales_matrix(sheet)
                rows.extend(sheet_rows)
        
        else:
            # Parse CSV file with matrix format
            import csv
            content = file_content.decode('utf-8-sig')
            csv_reader = csv.reader(io.StringIO(content))
            
            # Get headers (first row)
            headers = next(csv_reader)
            categories = []
            date_col_idx = 7  # Column H
            
            # Identify category columns (B-G, indices 1-6)
            for idx in range(1, 7):
                if idx < len(headers) and headers[idx]:
                    categories.append((idx, headers[idx].strip()))
            
            # If no categories found, use default
            if not categories:
                categories = [(i, f"Categoria_{i-1}") for i in range(1, 7)]
            
            # Parse data rows
            for row in csv_reader:
                if not any(row):  # Skip empty rows
                    continue
                
                # Column A: Sales channel
                canale_vendita = row[0].strip() if row[0] else ""
                if not canale_vendita:
                    continue
                
                # Column H: Date
                data_str = ""
                if len(row) > date_col_idx and row[date_col_idx]:
                    parsed_date = parse_date_string(row[date_col_idx])
                    if parsed_date:
                        data_str = parsed_date.strftime("%Y-%m-%d %H:%M:%S")
                
                if not data_str:
                    continue
                
                # Process each category column
                for col_idx, categoria_prodotto in categories:
                    if len(row) > col_idx and row[col_idx]:
                        quantita = parse_number(row[col_idx])
                        
                        if quantita > 0:
                            parsed_row = ParsedSalesRow(
                                data=data_str,
                                canale_vendita=canale_vendita,
                                categoria_prodotto=categoria_prodotto,
                                quantita=quantita,
                                prezzo_totale=None
                            )
                            rows.append(parsed_row)
        
        logger.info(f"Total parsed sales rows: {len(rows)}")
        return ParseSalesResponse(rows=rows, total_rows=len(rows))
    
    except Exception as e:
        logger.error(f"Error parsing sales file: {e}", exc_info=True)
        raise HTTPException(status_code=400, detail=f"Errore nel parsing del file: {str(e)}")


@router.post("/import-sales", response_model=ImportResponse)
async def import_sales_data(
    data: ImportSalesRequest,
    current_user: UserResponse = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Import parsed sales data into database"""
    try:
        vendite_service = VenditeService(db)
        
        imported_count = 0
        
        for row in data.rows:
            # Parse date
            data_dt = parse_date_to_datetime(row.data)
            
            if not data_dt:
                logger.warning(f"Skipping row with invalid date: {row.data}")
                continue
            
            # Create sales record
            await vendite_service.create({
                "data": data_dt,
                "canale_vendita": row.canale_vendita,
                "categoria_prodotto": row.categoria_prodotto,
                "quantita": row.quantita,
                "prezzo_totale": row.prezzo_totale,
                "created_at": datetime.now()
            }, user_id=current_user.id)
            
            imported_count += 1
        
        return ImportResponse(
            success=True,
            imported_count=imported_count,
            message=f"Importate con successo {imported_count} vendite"
        )
    
    except Exception as e:
        logger.error(f"Error importing sales data: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"Errore durante l'importazione: {str(e)}")